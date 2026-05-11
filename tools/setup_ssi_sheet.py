"""
setup_ssi_sheet.py — One-time setup: creates the Wandar SSI Tracker Google Sheet on Drive.

Run once:
  python tools/setup_ssi_sheet.py

What it does:
  1. Creates a Google Sheet named "Wandar SSI Tracker" with SSI column headers
  2. Saves the Sheet ID to tools/config.json as 'gdrive_ssi_sheet_id'
  3. Prints the URL — bookmark it and add a row each week before the Sunday pipeline

Requirements:
  - Google Drive credentials configured (same as upload_gdrive.py)
  - GDRIVE_TOKEN_JSON env var set (or tools/.gdrive-token.json exists locally)
"""

import sys
import json
import os
import tempfile
from pathlib import Path

TOOLS = Path(__file__).parent
CONFIG_PATH = TOOLS / "config.json"
SCOPES = ["https://www.googleapis.com/auth/drive.file"]

HEADERS = [
    "Date",
    "Total SSI",
    "Professional Brand",
    "Find the Right People",
    "Engage with Insights",
    "Build Relationships",
]


def _get_service():
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from google.auth.transport.requests import Request
    from googleapiclient.discovery import build

    TOKEN_PATH = TOOLS / ".gdrive-token.json"
    CREDENTIALS_PATH = TOOLS / "credentials.json"
    creds = None

    token_json_env = os.environ.get("GDRIVE_TOKEN_JSON")
    if token_json_env:
        creds = Credentials.from_authorized_user_info(json.loads(token_json_env), SCOPES)
    elif TOKEN_PATH.exists():
        creds = Credentials.from_authorized_user_file(str(TOKEN_PATH), SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            creds_json_env = os.environ.get("GDRIVE_CREDENTIALS_JSON")
            if creds_json_env:
                with tempfile.NamedTemporaryFile(mode="w", suffix=".json", delete=False) as f:
                    f.write(creds_json_env)
                    tmp = f.name
                try:
                    flow = InstalledAppFlow.from_client_secrets_file(tmp, SCOPES)
                finally:
                    Path(tmp).unlink(missing_ok=True)
            elif CREDENTIALS_PATH.exists():
                flow = InstalledAppFlow.from_client_secrets_file(str(CREDENTIALS_PATH), SCOPES)
            else:
                print("Error: Google Drive credentials not configured.")
                sys.exit(1)
            creds = flow.run_local_server(port=0)
            if not token_json_env:
                TOKEN_PATH.write_text(creds.to_json(), encoding="utf-8")

    return build("drive", "v3", credentials=creds)


def _make_xlsx():
    """Create a minimal xlsx with SSI headers and one example row."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("Error: openpyxl not installed. Run: pip install openpyxl")
        sys.exit(1)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SSI Tracker"

    ws.append(HEADERS)

    header_fill = PatternFill(start_color="1B4332", end_color="1B4332", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Example row showing the expected format (replace with real data)
    ws.append(["2026-05-10", 55, 14, 12, 15, 14])

    ws.column_dimensions["A"].width = 14
    for col in ["B", "C", "D", "E", "F"]:
        ws.column_dimensions[col].width = 24

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb.save(tmp.name)
    tmp.close()
    return tmp.name


def main():
    try:
        from googleapiclient.http import MediaFileUpload
    except ImportError:
        print("Error: google-api-python-client not installed. Run: pip install google-api-python-client")
        sys.exit(1)

    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8")) if CONFIG_PATH.exists() else {}
    existing_id = config.get("gdrive_ssi_sheet_id", "")

    if existing_id:
        url = f"https://docs.google.com/spreadsheets/d/{existing_id}/edit"
        print(f"SSI sheet already exists: {url}")
        print("To recreate: remove 'gdrive_ssi_sheet_id' from tools/config.json and re-run.")
        return

    print("Creating Wandar SSI Tracker Google Sheet...")
    xlsx_path = _make_xlsx()

    try:
        service = _get_service()
        mime_xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        mime_sheet = "application/vnd.google-apps.spreadsheet"

        media = MediaFileUpload(xlsx_path, mimetype=mime_xlsx, resumable=False)
        metadata = {"name": "Wandar SSI Tracker", "mimeType": mime_sheet}
        result = service.files().create(
            body=metadata,
            media_body=media,
            fields="id",
        ).execute()
        sheet_id = result["id"]
        media = None  # release file handle before cleanup
    finally:
        try:
            Path(xlsx_path).unlink(missing_ok=True)
        except PermissionError:
            pass  # Windows may keep the handle briefly; temp file will be cleaned up by OS

    config["gdrive_ssi_sheet_id"] = sheet_id
    CONFIG_PATH.write_text(json.dumps(config, indent=2), encoding="utf-8")

    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/edit"
    print(f"\n{'='*60}")
    print(f"SSI Tracker ready: {url}")
    print(f"{'='*60}")
    print(f"\nEach week before the Sunday pipeline, add one row:")
    print(f"  Date (YYYY-MM-DD) | Total /100 | Brand /25 | Find /25 | Engage /25 | Build /25")
    print(f"\nDelete the example row (row 2) before adding real data.")
    print(f"The pipeline reads the most recent row automatically.")


if __name__ == "__main__":
    main()
