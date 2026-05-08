"""
generate_word.py — Converts markdown post drafts or reports to branded .docx files.
Usage: python tools/generate_word.py <input_markdown_file>
Output: Same path with .docx extension.
"""

import sys
import re
from pathlib import Path
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("Error: python-docx not installed. Run: python -m pip install python-docx")
    sys.exit(1)

WANDAR_DARK = RGBColor(0x1A, 0x1A, 0x2E)
WANDAR_ACCENT = RGBColor(0xE8, 0xA0, 0x20)
WANDAR_GREY = RGBColor(0x6B, 0x6B, 0x6B)
BLACK = RGBColor(0x00, 0x00, 0x00)


def set_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    if level == 1:
        set_font(run, size=16, bold=True, color=WANDAR_DARK)
    elif level == 2:
        set_font(run, size=13, bold=True, color=WANDAR_DARK)
    else:
        set_font(run, size=11, bold=True, color=WANDAR_ACCENT)
    return p


def add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    label_run = p.add_run(f"{label}: ")
    set_font(label_run, size=10, bold=True, color=WANDAR_ACCENT)
    value_run = p.add_run(value)
    set_font(value_run, size=10, color=WANDAR_GREY)
    return p


def add_post_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    set_font(run, size=11, color=BLACK)
    p.paragraph_format.border_left = True
    return p


def add_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("─" * 60)
    set_font(run, size=9, color=WANDAR_GREY)
    return p


def parse_posts(text):
    posts = []
    blocks = re.split(r"\n---\n", text)
    current_post = {}
    body_lines = []

    def _save_current():
        if current_post:
            if body_lines and "body" not in current_post:
                current_post["body"] = "\n".join(body_lines).strip()
            if current_post.get("body") or current_post.get("pillar"):
                posts.append(dict(current_post))

    for block in blocks:
        block = block.strip()
        if not block:
            continue

        if re.match(r"PILLAR:|BUCKET:|FORMAT:|VIRALITY|BEST DAY:|POSTING WINDOW:", block):
            _save_current()
            current_post = {}
            body_lines = []
            for line in block.split("\n"):
                m = re.match(r"\*\*?([^:*]+)\*\*?:\s*(.+)", line)
                if m:
                    current_post[m.group(1).strip().lower().replace(" ", "_")] = m.group(2).strip()
                else:
                    m2 = re.match(r"([A-Z][A-Z &]+):\s*(.+)", line)
                    if m2:
                        current_post[m2.group(1).strip().lower().replace(" ", "_")] = m2.group(2).strip()

        elif block.startswith("WHY THIS WORKS:") or block.startswith("IMAGE DIRECTION:"):
            if body_lines:
                current_post["body"] = "\n".join(body_lines).strip()
                body_lines = []
            for line in block.split("\n"):
                if line.startswith("WHY THIS WORKS:"):
                    current_post["why_this_works"] = line.replace("WHY THIS WORKS:", "").strip()
                elif line.startswith("IMAGE DIRECTION:"):
                    current_post["image_direction"] = line.replace("IMAGE DIRECTION:", "").strip()

        else:
            # Skip markdown section headers and don't add body after WHY block is done
            if current_post and "why_this_works" not in current_post and not re.match(r"#+\s", block):
                body_lines.append(block)

    _save_current()
    return posts


def parse_report(text):
    return text


def build_post_doc(doc, posts, source_file):
    p = doc.add_paragraph()
    title_run = p.add_run("WANDAR — LinkedIn Posts")
    set_font(title_run, size=20, bold=True, color=WANDAR_DARK)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p2 = doc.add_paragraph()
    sub_run = p2.add_run(f"Generated: {datetime.now().strftime('%B %d, %Y')}  |  Source: {source_file}")
    set_font(sub_run, size=9, color=WANDAR_GREY)

    doc.add_paragraph()

    for idx, post in enumerate(posts, start=1):
        add_heading(doc, f"Post {idx} — {post.get('best_day', 'TBD')}", level=2)

        meta_fields = [
            ("Pillar", post.get("pillar", "")),
            ("Bucket", post.get("bucket", "")),
            ("Format", post.get("format", "")),
            ("Virality Strategy", post.get("virality_strategy", "")),
            ("Posting Window", post.get("posting_window", "")),
        ]
        for label, value in meta_fields:
            if value:
                add_label_value(doc, label, value)

        doc.add_paragraph()
        add_heading(doc, "Post Copy", level=3)

        body = post.get("body", "")
        if body:
            for para in body.split("\n\n"):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(para.strip())
                set_font(run, size=11)

        if post.get("why_this_works"):
            doc.add_paragraph()
            add_heading(doc, "Why This Works", level=3)
            p = doc.add_paragraph()
            run = p.add_run(post["why_this_works"])
            set_font(run, size=10, italic=True, color=WANDAR_GREY)

        if post.get("image_direction"):
            add_heading(doc, "Image Direction", level=3)
            p = doc.add_paragraph()
            run = p.add_run(post["image_direction"])
            set_font(run, size=10, italic=True, color=WANDAR_GREY)

        if idx < len(posts):
            add_divider(doc)
            doc.add_paragraph()


def build_report_doc(doc, text, source_file):
    p = doc.add_paragraph()
    title_run = p.add_run("WANDAR — Learning Loop Report")
    set_font(title_run, size=20, bold=True, color=WANDAR_DARK)

    p2 = doc.add_paragraph()
    sub_run = p2.add_run(f"Generated: {datetime.now().strftime('%B %d, %Y')}  |  Source: {source_file}")
    set_font(sub_run, size=9, color=WANDAR_GREY)

    doc.add_paragraph()

    for line in text.split("\n"):
        if line.startswith("# "):
            add_heading(doc, line[2:].strip(), level=1)
        elif line.startswith("## "):
            add_heading(doc, line[3:].strip(), level=2)
        elif line.startswith("### "):
            add_heading(doc, line[4:].strip(), level=3)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(line[2:].strip())
            set_font(run, size=11)
        elif line.strip():
            p = doc.add_paragraph()
            run = p.add_run(line.strip())
            set_font(run, size=11)


def main():
    if len(sys.argv) < 2:
        print("Usage: python tools/generate_word.py <input.md>")
        sys.exit(1)

    input_path = Path(sys.argv[1])
    if not input_path.exists():
        print(f"File not found: {input_path}")
        sys.exit(1)

    text = input_path.read_text(encoding="utf-8")
    output_path = input_path.with_suffix(".docx")

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    posts = parse_posts(text)

    if posts:
        build_post_doc(doc, posts, input_path.name)
    else:
        build_report_doc(doc, text, input_path.name)

    doc.save(output_path)
    print(f"Word document created: {output_path}")


if __name__ == "__main__":
    main()
