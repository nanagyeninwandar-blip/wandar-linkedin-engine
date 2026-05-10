"""
generate_analytics_template.py — Generates the Wandar LinkedIn analytics input Excel template.

Run once (called automatically by setup_analytics_sheet.py):
  python tools/generate_analytics_template.py

Output: analytics/analytics-tracker.xlsx
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent.parent
OUTPUT_PATH = ROOT / "analytics" / "analytics-tracker.xlsx"

WANDAR_DARK   = "1A1A2E"
WANDAR_ORANGE = "E8A020"
WANDAR_LIGHT  = "F5F5F0"
WHITE         = "FFFFFF"
GREEN         = "C6EFCE"
YELLOW        = "FFEB9C"
RED           = "FFC7CE"

HEADERS = [
    "Week",
    "Date",
    "Post Hook (first line)",
    "Impressions",
    "Saves",
    "Reposts",
    "Comments",
    "Operator Comments",
    "DMs Triggered",
    "Pillar",
    "Bucket",
    "Format",
]

PILLARS = [
    "1. How Safari Demand Forms",
    "2. Discovery & Distribution",
    "3. Trust & Conversion",
    "4. Social Listening",
    "5. Safari Buyer Intelligence",
    "6. Operator Growth & Positioning",
    "7. Founder Perspective",
    "8. Future of Safari Travel",
]

BUCKETS = ["Growth", "Authority", "Conversion", "Personal"]
FORMATS = ["insight", "story", "list", "contrast", "framework"]

COL_WIDTHS = {
    "Week": 14,
    "Date": 12,
    "Post Hook (first line)": 45,
    "Impressions": 13,
    "Saves": 10,
    "Reposts": 10,
    "Comments": 11,
    "Operator Comments": 18,
    "DMs Triggered": 14,
    "Pillar": 30,
    "Bucket": 14,
    "Format": 13,
}

THIN = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def make_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Weekly Analytics"
    ws.sheet_properties.tabColor = WANDAR_ORANGE

    # --- Instructions row ---
    ws.row_dimensions[1].height = 30
    instr = ws.cell(1, 1, "WANDAR LinkedIn Analytics Tracker — Fill in one row per post after each week. Run 'run learning loop' when done.")
    instr.font = Font(name="Calibri", size=10, color=WANDAR_DARK, italic=True)
    instr.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))

    # --- Header row ---
    ws.row_dimensions[2].height = 24
    header_fill = PatternFill("solid", fgColor=WANDAR_DARK)
    header_font = Font(name="Calibri", size=10, bold=True, color=WHITE)

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(2, col_idx, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
        ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS.get(header, 14)

    # --- Data rows (52 weeks × up to 5 posts = 260 rows pre-allocated) ---
    for row in range(3, 263):
        is_alt = (row % 2 == 0)
        row_fill = PatternFill("solid", fgColor=WANDAR_LIGHT if is_alt else WHITE)
        ws.row_dimensions[row].height = 18
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row, col)
            cell.fill = row_fill
            cell.font = Font(name="Calibri", size=10)
            cell.alignment = Alignment(vertical="center")
            cell.border = THIN

    # --- Data validation: Pillar dropdown ---
    pillar_col = HEADERS.index("Pillar") + 1
    dv_pillar = DataValidation(
        type="list",
        formula1='"' + ",".join(PILLARS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv_pillar.sqref = f"{get_column_letter(pillar_col)}3:{get_column_letter(pillar_col)}262"
    ws.add_data_validation(dv_pillar)

    # --- Data validation: Bucket dropdown ---
    bucket_col = HEADERS.index("Bucket") + 1
    dv_bucket = DataValidation(
        type="list",
        formula1='"' + ",".join(BUCKETS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv_bucket.sqref = f"{get_column_letter(bucket_col)}3:{get_column_letter(bucket_col)}262"
    ws.add_data_validation(dv_bucket)

    # --- Data validation: Format dropdown ---
    format_col = HEADERS.index("Format") + 1
    dv_format = DataValidation(
        type="list",
        formula1='"' + ",".join(FORMATS) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv_format.sqref = f"{get_column_letter(format_col)}3:{get_column_letter(format_col)}262"
    ws.add_data_validation(dv_format)

    # --- Conditional formatting: Saves column ---
    from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
    saves_col = get_column_letter(HEADERS.index("Saves") + 1)
    saves_range = f"{saves_col}3:{saves_col}262"
    ws.conditional_formatting.add(saves_range, CellIsRule(operator="greaterThanOrEqual", formula=["10"], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add(saves_range, CellIsRule(operator="between", formula=["5", "9"], fill=PatternFill("solid", fgColor=YELLOW)))
    ws.conditional_formatting.add(saves_range, CellIsRule(operator="lessThan", formula=["5"], fill=PatternFill("solid", fgColor=RED)))

    # --- Freeze panes below header ---
    ws.freeze_panes = "A3"

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Template created: {OUTPUT_PATH}")
    return OUTPUT_PATH


if __name__ == "__main__":
    make_template()
