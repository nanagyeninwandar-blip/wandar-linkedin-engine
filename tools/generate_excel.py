"""
generate_excel.py — Converts markdown tables to formatted .xlsx files.
Usage: python tools/generate_excel.py <input_markdown_file>
Output: Same path with .xlsx extension.
"""

import sys
import re
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Error: openpyxl not installed. Run: python -m pip install openpyxl")
    sys.exit(1)

WANDAR_DARK = "1A1A2E"
WANDAR_ACCENT = "E8A020"
WANDAR_LIGHT = "F5F5F0"
WHITE = "FFFFFF"
SCORE_HIGH = "C6EFCE"
SCORE_MID = "FFEB9C"
SCORE_LOW = "FFC7CE"


def parse_markdown_tables(text):
    tables = []
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith("|") and line.endswith("|"):
            header_line = line
            headers = [h.strip() for h in header_line.strip("|").split("|")]
            i += 1
            if i < len(lines) and re.match(r"^\|[-| :]+\|$", lines[i].strip()):
                i += 1
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    rows.append(row)
                    i += 1
                tables.append({"headers": headers, "rows": rows})
            else:
                i += 1
        else:
            i += 1
    return tables


def score_to_color(value):
    try:
        score = float(str(value).split("/")[0].strip())
        if score >= 7.5:
            return SCORE_HIGH
        elif score >= 5.0:
            return SCORE_MID
        else:
            return SCORE_LOW
    except (ValueError, AttributeError):
        return None


def write_table_to_sheet(ws, table, title=""):
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    row_offset = 1

    if title:
        ws.cell(row=row_offset, column=1, value=title).font = Font(
            name="Calibri", size=12, bold=True, color=WANDAR_DARK
        )
        row_offset += 2

    headers = table["headers"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_offset, column=col_idx, value=header)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=WANDAR_DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    row_offset += 1

    for r_idx, row in enumerate(table["rows"]):
        fill_color = WHITE if r_idx % 2 == 0 else WANDAR_LIGHT
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_offset + r_idx, column=col_idx, value=value)
            cell.font = Font(name="Calibri", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

            score_col = any(k in headers[col_idx - 1].lower() for k in ["score", "relevance"])
            if score_col:
                sc = score_to_color(value)
                cell.fill = PatternFill("solid", fgColor=sc if sc else fill_color)
            else:
                cell.fill = PatternFill("solid", fgColor=fill_color)

    for col_idx in range(1, len(headers) + 1):
        max_len = max(
            len(str(headers[col_idx - 1])),
            max((len(str(r[col_idx - 1])) if col_idx - 1 < len(r) else 0) for r in table["rows"]) if table["rows"] else 0,
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 4, 12), 45)

    ws.row_dimensions[row_offset - 1].height = 22


def main():
    if len(sys.argv) < 2:
        print("Usage: python tools/generate_excel.py <input.md>")
        sys.exit(1)

    input_path = Path(sys.argv[1])
    if not input_path.exists():
        print(f"File not found: {input_path}")
        sys.exit(1)

    text = input_path.read_text(encoding="utf-8")
    tables = parse_markdown_tables(text)

    if not tables:
        print(f"No markdown tables found in {input_path}")
        sys.exit(0)

    output_path = input_path.with_suffix(".xlsx")
    wb = Workbook()
    wb.remove(wb.active)

    section_names = re.findall(r"##\s+(.+)", text)

    for idx, table in enumerate(tables):
        sheet_name = section_names[idx] if idx < len(section_names) else f"Table {idx + 1}"
        sheet_name = sheet_name[:31]
        ws = wb.create_sheet(title=sheet_name)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = WANDAR_ACCENT
        write_table_to_sheet(ws, table, title=sheet_name)

    wb.save(output_path)
    print(f"Excel file created: {output_path}")


if __name__ == "__main__":
    main()
